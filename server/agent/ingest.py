"""Order documents (customer PO, SO ack, supplier PO) to structured facts.

1. Normalize the format in code: page text and/or page images, plus link annotations.
2. One structured-output LLM call that copies values verbatim with their location.
3. Code checks the result: every raw value must appear in the source text (grounding),
   and line arithmetic must tie out. One retry with feedback, then warnings, never guesses.
"""
from __future__ import annotations

import email
import email.policy
import hashlib
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf

from . import config, llm, parse
from .models import Extracted, OrderDoc, Val

IMAGE_TYPES = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".gif": "image/gif", ".webp": "image/webp"}


@dataclass
class Page:
    text: str | None = None
    image: bytes | None = None
    mime: str | None = None


@dataclass
class Normalized:
    name: str
    kind: str  # pdf_text | pdf_scan | image | xlsx | docx | email
    sha256: str
    pages: list[Page] = field(default_factory=list)
    links: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n".join(p.text for p in self.pages if p.text)

    @property
    def has_text(self) -> bool:
        return bool(self.text.strip())


# ---------------------------------------------------------------- normalization


def normalize_file(path: Path) -> Normalized:
    return normalize_bytes(path.name, path.read_bytes())


def normalize_bytes(name: str, data: bytes) -> Normalized:
    sha = hashlib.sha256(data).hexdigest()
    ext = Path(name).suffix.lower()
    if ext == ".pdf":
        return _pdf(name, data, sha)
    if ext in IMAGE_TYPES:
        return Normalized(name, "image", sha, [Page(image=data, mime=IMAGE_TYPES[ext])])
    if ext in (".xlsx", ".xlsm"):
        return Normalized(name, "xlsx", sha, [Page(text=_xlsx_text(data))])
    if ext == ".docx":
        return Normalized(name, "docx", sha, [Page(text=_docx_text(data))])
    if ext == ".eml":
        msg = email.message_from_bytes(data, policy=email.policy.default)
        body = msg.get_body(preferencelist=("plain", "html"))
        return normalize_email_body(name, msg["From"] or "", msg["Subject"] or "", body.get_content() if body else "", sha)
    raise ValueError(f"unsupported document type: {name}")


def normalize_email_body(name: str, sender: str, subject: str, body: str, sha: str | None = None) -> Normalized:
    # Graph can hand back quoted-printable soft breaks undecoded ("10/18=\n/2026").
    body = re.sub(r"=\r?\n", "", body)
    text = f"From: {sender}\nSubject: {subject}\n\n{body}"
    return Normalized(name, "email", sha or hashlib.sha256(text.encode()).hexdigest(), [Page(text=text)])


def _pdf(name: str, data: bytes, sha: str) -> Normalized:
    doc = pymupdf.open(stream=data, filetype="pdf")
    pages, links = [], []
    for page in doc:
        text = page.get_text("text", sort=True)
        image = page.get_pixmap(dpi=150).tobytes("png")
        pages.append(Page(text=text if text.strip() else None, image=image, mime="image/png"))
        links += [l["uri"] for l in page.get_links() if l.get("uri")]
    kind = "pdf_text" if any(p.text for p in pages) else "pdf_scan"
    return Normalized(name, kind, sha, pages, links)


def _cell(value) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _xlsx_text(data: bytes) -> str:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception:  # openpyxl rejects some real-world files (odd styles); read the cells directly
        return _xlsx_text_raw(data)
    out = []
    for ws in wb.worksheets:
        out.append(f"[sheet {ws.title}]")
        for row in ws.iter_rows():
            cells = [f"{c.coordinate}={_cell(c.value)}" for c in row if c.value not in (None, "")]
            if cells:
                out.append(" ; ".join(cells))
    return "\n".join(out)


def _xlsx_text_raw(data: bytes) -> str:
    """Cells straight from the sheet XML: shared strings, inline strings and values. No styles."""
    import zipfile
    from xml.etree import ElementTree as ET

    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    z = zipfile.ZipFile(io.BytesIO(data))
    shared = []
    if "xl/sharedStrings.xml" in z.namelist():
        for si in ET.fromstring(z.read("xl/sharedStrings.xml")).findall("m:si", ns):
            shared.append("".join(t.text or "" for t in si.iter(f"{{{ns['m']}}}t")))
    out = []
    for name in sorted(n for n in z.namelist() if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")):
        out.append(f"[{name.rsplit('/', 1)[-1]}]")
        for row in ET.fromstring(z.read(name)).iter(f"{{{ns['m']}}}row"):
            cells = []
            for c in row.findall("m:c", ns):
                v, t = c.find("m:v", ns), c.get("t")
                if t == "s" and v is not None:
                    text = shared[int(v.text)]
                elif t == "inlineStr":
                    text = "".join(x.text or "" for x in c.iter(f"{{{ns['m']}}}t"))
                else:
                    text = v.text if v is not None else ""
                if text not in ("", None):
                    cells.append(f"{c.get('r')}={text}")
            if cells:
                out.append(" ; ".join(cells))
    return "\n".join(out)


def _docx_text(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    out = [p.text for p in d.paragraphs if p.text.strip()]
    for i, table in enumerate(d.tables, 1):
        out.append(f"[table {i}]")
        for row in table.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)


# ---------------------------------------------------------------- extraction

INSTRUCTIONS = """You read order documents for Demo Distributor, Inc., an electronic component distributor.
The document is one of: a customer purchase order sent to Demo Distributor (customer_po), Demo Distributor's
sales order acknowledgement (so_ack), Demo Distributor's purchase order to a supplier (supplier_po), a supplier's
acknowledgement or order confirmation of Demo Distributor's purchase order (supplier_ack), or other.

Fill the schema by COPYING values exactly as printed. Do not reformat, round, translate, normalise or correct
anything: keep thousands separators, decimal commas, currency symbols, leading zeros, spaces and punctuation in
part numbers and account numbers. If a value is not on the document, use null. Never infer a value that is not
printed. For every value give its location: page number (1-based) and `where`, a short anchor such as the field
label, column name, line number or spreadsheet cell reference.

Field guidance:
- doc_number: the PO number (customer_po, supplier_po), the sales order / order number (so_ack), or the supplier's
  own order or confirmation number (supplier_ack).
- customer_po_ref: on an acknowledgement, the buyer's PO number it confirms ("Cust PO #", "Your PO", "PO Ref").
  On a supplier_ack that is Demo Distributor's PO number. On a supplier_po, the customer SO it is bought for, if printed.
- customer_name / customer_id: the buying company and any customer id printed (e.g. "Customer Id C1002").
- addressed_to: the vendor/supplier block on a PO (who the PO is sent to).
- bill_to / ship_to: the complete address block (company, attention or contact line, street, city, region,
  postcode, country), one printed line per line, separated by newlines. If the document names a delivery
  address only in a note or email text, use that.
- carrier_account: a shipping/freight account number, e.g. from "FedEx Ground on our account 000601844" or
  "Customer Freight Acct # - 000601844". Copy only the number as printed.
- ship_via: the carrier/service text. incoterm: FOB / Incoterms / shipping terms text (e.g. "Ex Works",
  "FOB Destination", "FCA Austin"). freight_terms: freight payment terms on an ack (e.g. "Prepay & Add").
- ship_option: an explicit ship-complete / partial-shipment instruction or field (e.g. "Ship Partials",
  "Ship Complete", or the sentence "Ship complete. No partial shipments without written approval.").
- lines: one entry per line item, in document order. customer_pn is the buyer's own item number (e.g.
  "MER-61495", "Cust Item No : COP-42702" -> "COP-42702"); mpn is the manufacturer part number. Buyers often
  print their own item code in the Item or Description column, above or instead of a description (e.g. a bold
  "05673-001-XTD" where other lines have a description): that code is customer_pn, not description.
  On a supplier_ack, ship_date is the supplier's promised or confirmed ship/delivery date. country_of_origin and
  date_code are per-line values when printed (e.g. "COO: CN", "Date code 2414"). On an ack,
  ship_date is "Est Ship" and required_date is "Req Dt". line_notes holds text printed under a line item.
- notes: every free-text note, condition or instruction paragraph, verbatim, one item per paragraph or numbered
  clause (including conditions on later pages).
- references: references to terms and conditions, clause codes, document ids or links (verbatim).
Text extracted from the file is provided when available, together with page images. Prefer the extracted text
for exact characters; use the image for layout and for anything not in the text."""


def repair(x: Extracted) -> Extracted:
    """Code fixes for known extraction slips, applied to every extraction (fresh, cached or saved).

    A line with no manufacturer part number but a 'customer part number' and a manufacturer: the
    model filed the MPN under the wrong field (seen on a photographed PO with a 'Part Number'
    column), so the number moves to mpn. A customer part number identical to the MPN is dropped."""
    for i, l in enumerate(x.doc.lines, 1):
        if not l.mpn.raw and l.customer_pn.raw and l.manufacturer.raw:
            l.mpn, l.customer_pn = l.customer_pn, Val(raw=None, loc=None)
            note = f"line {i}: part number {l.mpn.raw!r} was extracted as a customer part number; moved to MPN"
            if note not in x.warnings:
                x.warnings.append(note)
        elif l.customer_pn.raw and parse.mpn_key(l.customer_pn.raw) == parse.mpn_key(l.mpn.raw):
            l.customer_pn = Val(raw=None, loc=None)
    return x


def extract(norm: Normalized, use_cache: bool = True) -> Extracted:
    prompt_id = hashlib.sha256(INSTRUCTIONS.encode()).hexdigest()[:8]
    cache = config.STORE_DIR / "cache" / "extract" / f"{norm.sha256}.{config.EXTRACT_MODEL}.{prompt_id}.json"
    if use_cache and cache.exists():
        return repair(Extracted.model_validate_json(cache.read_text()))

    content = _content(norm)
    doc = llm.parse(config.EXTRACT_MODEL, INSTRUCTIONS, content, OrderDoc)
    problems = check_extraction(doc, norm)
    if problems:
        feedback = "Your previous extraction had these problems. Re-read the document and fix them:\n- " + "\n- ".join(problems)
        doc = llm.parse(config.EXTRACT_MODEL, INSTRUCTIONS, content + [llm.text_part(feedback)], OrderDoc)
        problems = check_extraction(doc, norm)

    method = "text+vision" if norm.has_text and any(p.image for p in norm.pages) else ("text" if norm.has_text else "vision")
    result = Extracted(doc=doc, source=norm.name, sha256=norm.sha256, method=method, links=norm.links, warnings=problems)
    cache.parent.mkdir(parents=True, exist_ok=True)
    cache.write_text(result.model_dump_json(indent=1))
    return repair(result)


def _content(norm: Normalized) -> list[dict]:
    parts = [llm.text_part(f"Document: {norm.name} (format: {norm.kind})")]
    for i, page in enumerate(norm.pages, 1):
        if page.text:
            parts.append(llm.text_part(f"--- page {i} extracted text ---\n{page.text}"))
        if page.image:
            parts.append(llm.text_part(f"--- page {i} image ---"))
            parts.append(llm.image_part(page.image, page.mime or "image/png"))
    if norm.links:
        parts.append(llm.text_part("Link targets found in the file: " + ", ".join(norm.links)))
    return parts


def iter_vals(doc: OrderDoc):
    """Every Val in the document with a readable path, for grounding checks."""
    for name, v in doc.header:
        if isinstance(v, Val):
            yield f"header.{name}", v
    for i, line in enumerate(doc.lines, 1):
        for name, v in line:
            if isinstance(v, Val):
                yield f"line {i}.{name}", v
            elif isinstance(v, list):
                for j, n in enumerate(v):
                    yield f"line {i}.{name}[{j}]", n
    for name, v in doc.totals:
        yield f"totals.{name}", v
    for i, v in enumerate(doc.notes):
        yield f"notes[{i}]", v
    for i, v in enumerate(doc.references):
        yield f"references[{i}]", v


def check_extraction(doc: OrderDoc, norm: Normalized) -> list[str]:
    """Problems that justify a retry: ungrounded values (text sources) and lines that don't tie out."""
    problems = []
    if norm.has_text:
        source = parse.squash(norm.text)
        for path, v in iter_vals(doc):
            # Multi-line values (address blocks) are checked line by line: columns that sit side
            # by side on the page interleave in the text layer.
            parts = [p for p in (v.raw or "").splitlines() if p.strip()]
            missing = [p for p in parts if parse.squash(p) not in source]
            if missing:
                problems.append(f"{path} = {v.raw!r} does not appear verbatim in the document text")
    raws = [(l.qty.raw, l.unit_price.raw, l.extended.raw) for l in doc.lines]
    comma = parse.detect_decimal_comma(raws)
    for i, (q, p, e) in enumerate(raws, 1):
        if parse.ties_out(parse.num(q, comma), parse.num(p, comma), parse.num(e, comma)) is False:
            problems.append(f"line {i}: quantity {q!r} x unit price {p!r} does not equal extended {e!r}")
    return problems
